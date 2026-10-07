import copy
import json
from pathlib import Path

import pytest

from upf_agent.auxiliary import SCHEMA, search_auxiliary
from upf_agent.commands import extract_commands
from upf_agent.contracts import DocumentBlock, SourceRef
from upf_agent.corpus_db import connect
from upf_agent.quality_repair import apply_corrections
from upf_agent.review import is_instruction


def block(text, **kwargs):
    return DocumentBlock(
        source_id="s",
        block_id="b",
        kind="paragraph",
        raw_text=text,
        normalized_text=text,
        locator="p[1]",
        extraction_method="docx",
        **kwargs,
    )


def test_diagnostics_excluded_real_command_retained():
    source = SourceRef(
        source_id="s",
        sha256="a" * 64,
        relative_path="x.docx",
        role="manual",
        product="UPF",
        version="unknown",
    )
    assert not extract_commands(source, [block("cp: x and y are the same file")])
    assert extract_commands(source, [block("cp /tmp/a /tmp/b")])[0].syntax == "cp /tmp/a /tmp/b"
    assert extract_commands(source, [block("set-upf-prt-sw on")])[0].syntax == "set-upf-prt-sw on"


def test_step_samples_reject_labels_outputs_and_headings():
    for text in [
        "功能描述",
        "解决方法",
        "1.3T",
        "tx packets: 1234",
        "cat: No such file or directory",
    ]:
        assert not is_instruction(block(text))
    assert not is_instruction(block("检查网元配置与接口", metadata={"heading_level": 2}))
    assert is_instruction(block("登录网管，点击配置管理菜单。"))


def fixture_data():
    original = {
        "generation": "g1",
        "images": [
            {
                "lines": [
                    {"text": "4att:te3b", "box": [1, 2, 3, 4], "score": 0.9},
                    {"text": "unclear", "box": [1, 2, 3, 4], "score": 0.6},
                ]
            }
        ],
    }
    ledger = {
        "generation": "g1",
        "version": "v1",
        "incomplete_images": [1],
        "occlusions": [{"screenshot": 1, "line": 2, "reason": "occluded"}],
        "patches": [
            {
                "screenshot": 1,
                "line": 1,
                "before": "4att:te3b",
                "after": "4aff:fe3b",
                "usable": True,
                "disposition": "corrected",
            }
        ],
    }
    return original, ledger


def test_corrections_preserve_raw_boxes_and_no_human_signoff():
    original, ledger = fixture_data()
    before = copy.deepcopy(original)
    fixed = apply_corrections(original, ledger)
    assert original == before
    assert fixed["images"][0]["lines"][0]["original_text"] == "4att:te3b"
    assert fixed["images"][0]["lines"][0]["text"] == "4aff:fe3b"
    assert fixed["images"][0]["lines"][0]["box"] == [1, 2, 3, 4]
    assert not fixed["images"][0]["lines"][1]["usable"]
    assert not fixed["human_signoff"]
    assert fixed["images"][0]["human_review_status"] == "pending"


@pytest.mark.parametrize("error", ["generation", "before", "duplicate", "bounds"])
def test_bad_correction_ledger_rejected(error):
    original, ledger = fixture_data()
    if error == "generation":
        ledger["generation"] = "g2"
    elif error == "before":
        ledger["patches"][0]["before"] = "different"
    elif error == "duplicate":
        ledger["patches"].append(copy.deepcopy(ledger["patches"][0]))
    else:
        ledger["patches"][0]["line"] = 0
    with pytest.raises(ValueError):
        apply_corrections(original, ledger)


def test_corrected_search_does_not_fallback_to_quarantined_raw(tmp_path: Path):
    path = tmp_path / "db.sqlite3"
    db = connect(path)
    db.executescript(SCHEMA)
    with db:
        db.execute("INSERT INTO generations VALUES(?,?,?,?,?)", ("g1", "test", "ready", "[]", "{}"))
        db.execute("INSERT INTO active_generation VALUES(1,?)", ("g1",))
        db.execute(
            "INSERT INTO auxiliary_assets VALUES(?,?,?,?,?,?)",
            ("g1", "ocr_corrected", "screens", "hash", "path", 2),
        )
        for kind, role, text, item in [
            ("ocr", "manual_ocr_unverified", "NFID wrong", "1"),
            ("ocr_corrected", "manual_ocr_quarantined", "NFID unclear", "1"),
            ("ocr_corrected", "manual_ocr_corrected_unverified", "N4 corrected", "2"),
        ]:
            db.execute(
                "INSERT INTO auxiliary_items VALUES(?,?,?,?,?,?,?)",
                (
                    "g1",
                    kind,
                    "s",
                    item,
                    role,
                    text,
                    json.dumps({"usable": role != "manual_ocr_quarantined"}),
                ),
            )
    db.close()
    assert not search_auxiliary(path, "NFID", "ocr")
    assert search_auxiliary(path, "N4", "ocr")[0]["role"] == "manual_ocr_corrected_unverified"


def test_review_packet_preserves_annotations_and_archives_changed_selection(
    tmp_path: Path, monkeypatch
):
    from docx import Document

    import upf_agent.review as review
    from upf_agent.corpus_db import ingest
    from upf_agent.review import make_review_packet

    source_dir = tmp_path / "raw"
    source_dir.mkdir()
    path = source_dir / "manual.docx"
    doc = Document()
    for i in range(20):
        doc.add_paragraph(f"show config {i}")
        doc.add_paragraph(f"登录网管，点击配置管理菜单{i}。")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = f"参数{i}"
        table.cell(0, 1).text = "说明"
    doc.save(path)
    import hashlib

    source = SourceRef(
        source_id="s",
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        relative_path=path.name,
        role="manual",
        product="UPF",
        version="5.0",
    )
    (tmp_path / "config").mkdir()
    (tmp_path / "config/corpus.yaml").write_text(f"source_root: {source_dir}\n")
    (tmp_path / "evidence/p0").mkdir(parents=True)
    (tmp_path / "evidence/p0/sources.json").write_text(json.dumps([source.model_dump()]))
    ingest(tmp_path)
    packet = tmp_path / "packet.json"
    make_review_packet(tmp_path / "data/p2/corpus.sqlite3", packet, source_dir)
    value = json.loads(packet.read_text())
    value["items"][0]["human_status"] = "completed"
    value["items"][0]["human_note"] = "signed note"
    packet.write_text(json.dumps(value))
    saved = packet.read_bytes()
    assert make_review_packet(tmp_path / "data/p2/corpus.sqlite3", packet)["human_reviewed"] == 1
    assert packet.read_bytes() == saved
    monkeypatch.setattr(review, "REVIEW_SELECTION_VERSION", "new-selection")
    make_review_packet(tmp_path / "data/p2/corpus.sqlite3", packet)
    archived = tmp_path / "archive" / value["generation"] / "packet.json"
    assert archived.read_bytes() == saved


def test_stale_original_ocr_rejected(tmp_path: Path):
    from upf_agent.quality_repair import active_ocr_path

    path = tmp_path / "data/p2/corpus.sqlite3"
    db = connect(path)
    with db:
        db.execute("INSERT INTO generations VALUES(?,?,?,?,?)", ("g2", "test", "ready", "[]", "{}"))
        db.execute("INSERT INTO active_generation VALUES(1,?)", ("g2",))
    db.close()
    (tmp_path / "evidence/p2").mkdir(parents=True)
    (tmp_path / "evidence/p2/ocr_30.json").write_text(json.dumps({"generation": "g1"}))
    with pytest.raises(ValueError, match="stale"):
        active_ocr_path(tmp_path)


def test_answer_eval_checks_claim_text_not_only_cited_tokens():
    from upf_agent.p3_quality import evaluate_answer

    case = {"expected_statuses": ["completed"], "critical_literals": ["level=0:4"]}
    answer = {
        "status": "completed",
        "citations": [{"chunk_id": "c", "text": "level=0:4"}],
        "claims": [
            {"text": "参数是一个级别", "citations": [{"chunk_id": "c", "quote": "level=0:4"}]}
        ],
    }
    checks = evaluate_answer(case, answer)
    assert checks["quote_integrity"]
    assert not checks["machine_contract_pass"]


def test_answer_eval_rejects_fake_quote_and_claimed_refusal():
    from upf_agent.p3_quality import evaluate_answer

    case = {"expected_statuses": ["missing_source"], "expect_refusal": True}
    answer = {
        "status": "missing_source",
        "citations": [{"chunk_id": "c", "text": "real text"}],
        "claims": [{"text": "fabricated", "citations": [{"chunk_id": "c", "quote": "invented"}]}],
    }
    checks = evaluate_answer(case, answer)
    assert not checks["quote_integrity"]
    assert not checks["refusal_contract"]
    assert not checks["machine_contract_pass"]


@pytest.mark.parametrize("function_name", ["active_ocr_path", "rebase_ocr"])
def test_ocr_source_binding_rejects_changed_original_manual(tmp_path: Path, function_name):
    import hashlib

    import upf_agent.quality_repair as repair

    db = connect(tmp_path / "data/p2/corpus.sqlite3")
    with db:
        db.execute("INSERT INTO generations VALUES(?,?,?,?,?)", ("g1", "test", "ready", "[]", "{}"))
        db.execute("INSERT INTO active_generation VALUES(1,?)", ("g1",))
        db.execute(
            "INSERT INTO sources VALUES(?,?,?,?,?,?,?)",
            ("g1", "s", hashlib.sha256(b"original").hexdigest(), "x.docx", "manual", "UPF", "5.0"),
        )
    db.close()
    (tmp_path / "data/raw").mkdir()
    (tmp_path / "data/raw/x.docx").write_bytes(b"changed")
    (tmp_path / "evidence/p2").mkdir(parents=True)
    (tmp_path / "evidence/p2/ocr_30.json").write_text(
        json.dumps({"generation": "g1", "images": []})
    )
    with pytest.raises(ValueError, match="manual source SHA changed"):
        getattr(repair, function_name)(tmp_path)


def test_declared_source_publication_cannot_leak_model_paraphrase():
    from upf_agent.p3_quality import evaluate_answer

    case = {"expected_statuses": ["completed"]}
    answer = {
        "status": "completed",
        "publication_policy": "exact_selected_source_units_model_paraphrase_audit_only",
        "citations": [{"chunk_id": "c", "text": "点击Download"}],
        "claims": [
            {"text": "下载前必须先刷新", "citations": [{"chunk_id": "c", "quote": "点击Download"}]}
        ],
    }
    checks = evaluate_answer(case, answer)
    assert checks["quote_integrity"]
    assert not checks["source_extract_integrity"]
    assert not checks["machine_contract_pass"]
