import hashlib
import json
from pathlib import Path

from docx import Document

from upf_agent.commands import extract_commands
from upf_agent.contracts import SourceRef
from upf_agent.corpus_db import active_id, connect, ingest, query_blocks
from upf_agent.docx_parse import DocxParser
from upf_agent.review import _source_text


def _source(path: Path) -> SourceRef:
    return SourceRef(
        source_id="manual-test",
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        relative_path=path.name,
        role="manual",
        product="UPF",
        version="5.0",
    )


def test_docx_order_table_and_exact_locator(tmp_path: Path):
    path = tmp_path / "tiny.docx"
    doc = Document()
    doc.add_heading("检查 N4", level=1)
    doc.add_paragraph("show cp-node")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "参数"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "n4_address"
    table.cell(1, 1).text = "N4地址"
    doc.add_paragraph("完成检查")
    doc.save(path)
    source = _source(path)
    parsed = DocxParser(source, path, tmp_path / "media").parse()
    assert [b.kind for b in parsed.blocks if b.parent_id is None] == [
        "paragraph",
        "paragraph",
        "table",
        "paragraph",
    ]
    command = extract_commands(source, parsed.blocks)
    assert len(command) == 1 and command[0].syntax == "show cp-node"
    for block in parsed.blocks:
        if block.kind in {"paragraph", "table_row"}:
            assert _source_text(path, block.locator) == block.raw_text


def test_ingest_idempotent_and_scoped_to_version(tmp_path: Path):
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    path = source_dir / "tiny.docx"
    doc = Document()
    doc.add_paragraph("show config")
    doc.save(path)
    source = _source(path)
    project = tmp_path / "project"
    (project / "config").mkdir(parents=True)
    (project / "evidence/p0").mkdir(parents=True)
    (project / "config/corpus.yaml").write_text(f"source_root: {source_dir}\n", encoding="utf-8")
    (project / "evidence/p0/sources.json").write_text(
        json.dumps([source.model_dump()]), encoding="utf-8"
    )
    first = ingest(project)
    second = ingest(project)
    assert first["generation"] == second["generation"]
    assert not first["reused"] and second["reused"]
    assert first["stats"]["commands"] == 1
    assert len(query_blocks(project / "data/p2/corpus.sqlite3", product="UPF", version="5.0")) == 1
    assert not query_blocks(project / "data/p2/corpus.sqlite3", product="UPF", version="4.0")
    changed = ingest(project, parser_version="forced-rebuild")
    assert changed["generation"] != first["generation"]
    db = connect(project / "data/p2/corpus.sqlite3")
    try:
        assert active_id(db) == changed["generation"]
    finally:
        db.close()


def test_ocr_never_becomes_command_authority():
    from upf_agent.contracts import DocumentBlock

    source = SourceRef(
        source_id="s",
        sha256="a" * 64,
        relative_path="x.docx",
        role="manual",
        product="UPF",
        version="unknown",
    )
    block = DocumentBlock(
        source_id="s",
        block_id="b",
        kind="image_ocr",
        raw_text="show config",
        normalized_text="show config",
        locator="image[1]",
        extraction_method="ocr",
        quality_flags=["ocr_unverified"],
    )
    assert extract_commands(source, [block]) == []


def test_metric_labels_are_not_commands():
    from upf_agent.contracts import DocumentBlock

    source = SourceRef(
        source_id="s",
        sha256="a" * 64,
        relative_path="x.docx",
        role="manual",
        product="UPF",
        version="unknown",
    )
    labels = ["Service Request Success", "Delete SM Policy Success Number", "IP Version"]
    blocks = [
        DocumentBlock(
            source_id="s",
            block_id=str(i),
            kind="paragraph",
            raw_text=label,
            normalized_text=label,
            locator=f"p[{i}]",
            extraction_method="test",
        )
        for i, label in enumerate(labels)
    ]
    assert extract_commands(source, blocks) == []


def test_placeholder_command_keeps_original_characters():
    from upf_agent.contracts import DocumentBlock

    source = SourceRef(
        source_id="s",
        sha256="a" * 64,
        relative_path="x.docx",
        role="manual",
        product="UPF",
        version="unknown",
    )
    raw = "create gre tunnel src {UPF1的N6地址} dst {UPF2的N6地址}"
    block = DocumentBlock(
        source_id="s",
        block_id="b",
        kind="paragraph",
        raw_text=raw,
        normalized_text=raw,
        locator="p[1]",
        extraction_method="test",
    )
    command = extract_commands(source, [block])[0]
    assert command.syntax == raw
    assert command.parameter_tokens == ["{UPF1的N6地址}", "{UPF2的N6地址}"]


def test_ocr_result_requires_matching_generation_and_image_hash(tmp_path: Path):
    from upf_agent.image_queue import record_ocr_results

    queue_path = tmp_path / "queue.json"
    result_path = tmp_path / "result.json"
    queue_path.write_text(
        json.dumps(
            {
                "generation": "g1",
                "selected": [{"block_id": "b1", "sha256": "a" * 64, "ocr_status": "pending"}],
            }
        )
    )
    result_path.write_text(
        json.dumps(
            {
                "generation": "g1",
                "images": [
                    {
                        "block_id": "b1",
                        "image_sha256": "a" * 64,
                        "lines": [{"text": "show config", "score": 0.9, "box": [1, 2, 3, 4]}],
                    }
                ],
            }
        )
    )
    assert record_ocr_results(queue_path, result_path)["ocr_lines"] == 1
    assert json.loads(queue_path.read_text())["selected"][0]["ocr_status"] == "completed_unverified"
    bad = json.loads(result_path.read_text())
    bad["images"][0]["image_sha256"] = "b" * 64
    result_path.write_text(json.dumps(bad))
    import pytest

    with pytest.raises(ValueError, match="SHA mismatch"):
        record_ocr_results(queue_path, result_path)


def test_auxiliary_import_is_idempotent_and_paper_is_isolated(tmp_path: Path):
    import pytest

    from upf_agent.auxiliary import attach_assets, search_auxiliary

    db_path = tmp_path / "corpus.sqlite3"
    db = connect(db_path)
    with db:
        db.execute("INSERT INTO generations VALUES(?,?,?,?,?)", ("g1", "test", "ready", "[]", "{}"))
        db.execute("INSERT INTO active_generation VALUES(1,?)", ("g1",))
    db.close()
    ocr_path = tmp_path / "ocr.json"
    pdf_path = tmp_path / "pdf.json"
    ocr_path.write_text(
        json.dumps(
            {
                "generation": "g1",
                "images": [
                    {
                        "source_id": "manual-1",
                        "block_id": "image-1",
                        "image_sha256": "a" * 64,
                        "locator": "image[1]",
                        "status": "ocr_unverified",
                        "lines": [{"text": "N4 address", "box": [0, 0, 2, 2], "score": 0.9}],
                    }
                ],
            }
        )
    )
    pdf_path.write_text(
        json.dumps(
            {
                "role": "paper",
                "source_id": "paper-1",
                "source_sha256": "b" * 64,
                "items": [
                    {"text": "research paper", "provenance": [{"page_no": 1}], "label": "text"}
                ],
            }
        )
    )
    first = attach_assets(db_path, ocr_path, pdf_path, "b" * 64)
    second = attach_assets(db_path, ocr_path, pdf_path, "b" * 64)
    assert not any(x["reused"] for x in first["assets"])
    assert all(x["reused"] for x in second["assets"])
    assert search_auxiliary(db_path, "N4", "ocr")[0]["role"] == "manual_ocr_unverified"
    assert search_auxiliary(db_path, "research", "pdf")[0]["role"] == "paper"
    assert not query_blocks(db_path, product="UPF")
    with pytest.raises(ValueError, match="Paper source role or SHA mismatch"):
        attach_assets(db_path, ocr_path, pdf_path, "c" * 64)
