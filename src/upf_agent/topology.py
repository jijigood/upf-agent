"""Bounded CSV/XLSX topology import with explicit aliases and source provenance."""

import csv
import hashlib
import json
from collections.abc import Iterator
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from pydantic import ValidationError

from upf_agent.contracts import ResultStatus, ToolResult, TopologyEdge
from upf_agent.log_parser import file_sha
from upf_agent.log_tools import Row, write_json

TOPOLOGY_VERSION = "topology-v1"
HEADERS = {
    "网元": "device_id",
    "端口": "device_port",
    "对端网元": "peer_device_id",
    "对端端口": "peer_port",
    "接口": "interface_type",
    "生效时间": "valid_from",
    "失效时间": "valid_to",
    "VLAN": "vlan",
    "用途": "purpose",
}
FIELDS = {
    "device_id",
    "device_port",
    "peer_device_id",
    "peer_port",
    "interface_type",
    "valid_from",
    "valid_to",
    "vlan",
    "purpose",
}


def canonical(value: str, aliases: dict[str, str]) -> str:
    visited: set[str] = set()
    while value in aliases and aliases[value] != value:
        if value in visited:
            raise ValueError("cyclic device aliases")
        visited.add(value)
        value = aliases[value]
    return value


def _rows(source: Path) -> Iterator[tuple[str | None, int, list[str], list[str], list[str]]]:
    if source.stat().st_size > 20 * 1024 * 1024:
        raise ValueError("topology files are bounded to 20MiB; split larger files")
    if source.suffix.lower() == ".csv":
        with source.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            headers = next(reader, [])
            for row in reader:
                yield None, reader.line_num, headers, row, []
    elif source.suffix.lower() == ".xlsx":
        # Limit expanded ZIP members before openpyxl allocates workbook cells.
        from zipfile import ZipFile

        with ZipFile(source) as archive:
            if sum(member.file_size for member in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("expanded XLSX exceeds 100MiB bound")
        book = load_workbook(source, data_only=True, keep_links=False)
        try:
            for sheet in book.worksheets:
                if sheet.max_row > 100000 or sheet.max_column > 100:
                    raise ValueError("topology sheet exceeds 100000 rows/100 columns")
                headers = [str(cell.value or "") for cell in sheet[1]]
                for number, cells in enumerate(sheet.iter_rows(min_row=2), 2):
                    values, flags = [], []
                    for cell in cells:
                        value = cell.value
                        if value is None:
                            for merged in sheet.merged_cells.ranges:
                                if cell.coordinate in merged:
                                    anchor = sheet.cell(merged.min_row, merged.min_col)
                                    value = anchor.value
                                    flags.append(f"merged_origin:{sheet.title}!{anchor.coordinate}")
                                    break
                        values.append(str(value) if value is not None else "")
                    yield sheet.title, number, headers, values, flags
        finally:
            book.close()
    else:
        raise ValueError("topology must be CSV or XLSX")


def import_topology(
    source: Path,
    root: Path,
    synthetic: bool = False,
    device_aliases: dict[str, str] | None = None,
    header_mapping: dict[str, str] | None = None,
    interface_aliases: dict[str, str] | None = None,
) -> Row:
    sha = file_sha(source)
    source_id = f"topology-{sha[:16]}"
    aliases, interfaces = device_aliases or {}, interface_aliases or {}
    for alias in aliases:
        canonical(alias, aliases)
    headers_map = {**HEADERS, **(header_mapping or {})}
    if any(value not in FIELDS for value in headers_map.values()):
        raise ValueError("header mapping uses unsupported fields")
    identity = {
        "source_sha256": sha,
        "version": TOPOLOGY_VERSION,
        "code_sha256": file_sha(Path(__file__)),
        "synthetic": synthetic,
        "device_aliases": aliases,
        "header_mapping": headers_map,
        "interface_aliases": interfaces,
    }
    generation = hashlib.sha256(json.dumps(identity, sort_keys=True).encode()).hexdigest()[:20]
    edges, quarantine = [], []
    for index, (sheet, number, headers, cells, flags) in enumerate(_rows(source)):
        if index >= 100000:
            raise ValueError("topology exceeds 100000 data rows")
        if not any(value.strip() for value in cells):
            continue
        mapped = [headers_map.get(value.strip(), value.strip()) for value in headers]
        data = {
            key: value.strip() for key, value in zip(mapped, cells, strict=False) if key in FIELDS
        }
        try:
            if len(set(mapped)) != len(mapped):
                raise ValueError("duplicate mapped headers")
            if len(cells) != len(headers):
                raise ValueError("row width differs from header")
            required = ["device_id", "device_port", "peer_device_id", "peer_port"]
            if any(not data.get(key) for key in required):
                raise ValueError("missing endpoint identity/port; no implicit forward fill")
            for key in ("device_id", "peer_device_id"):
                data[key] = canonical(data[key], aliases)
            interface = data.get("interface_type")
            interface = interfaces.get(interface, interface) if interface else None
            if interface and interface not in {"N3", "N4", "N6", "N9"}:
                flags.append("unrecognized_interface_preserved_exactly")
            start = datetime.fromisoformat(data["valid_from"]) if data.get("valid_from") else None
            end = datetime.fromisoformat(data["valid_to"]) if data.get("valid_to") else None
            if start is None and end is None:
                flags.append("validity_unknown")
            if start is None:
                flags.append("valid_from_unknown")
            if end is None:
                flags.append("valid_to_unknown")
            edge = TopologyEdge(
                edge_id=f"{source_id}:{sheet or 'csv'}:{number}",
                **{key: data[key] for key in required},
                interface_type=interface,
                valid_from=start,
                valid_to=end,
                source_id=source_id,
                source_sha256=sha,
                source_sheet=sheet,
                source_row=number,
                synthetic=synthetic,
                vlan=data.get("vlan"),
                purpose=data.get("purpose"),
                quality_flags=flags,
            )
            edges.append(edge.model_dump(mode="json"))
        except (ValueError, ValidationError) as error:
            quarantine.append(
                {
                    "source_sheet": sheet,
                    "source_row": number,
                    "raw_cells": cells,
                    "reason": str(error),
                }
            )
    if file_sha(source) != sha:
        raise ValueError("topology source changed during import")
    directory = root / generation
    directory.mkdir(parents=True, exist_ok=True)
    content = json.dumps(edges, ensure_ascii=False, indent=2) + "\n"
    target = directory / "edges.json"
    if target.exists() and target.read_text() != content:
        raise ValueError("immutable topology generation conflict")
    target.write_text(content, encoding="utf-8")
    write_json(directory / "quarantine.json", quarantine)
    manifest: Row = {
        **identity,
        "generation": generation,
        "source_path": str(source.resolve()),
        "source_id": source_id,
        "accepted_edges": len(edges),
        "quarantined_rows": len(quarantine),
        "edges_sha256": file_sha(target),
        "quarantine_sha256": file_sha(directory / "quarantine.json"),
    }
    write_json(directory / "manifest.json", manifest)
    return manifest


def lookup_topology(
    directory: Path, device_id: str, interface_type: str | None = None, at: datetime | None = None
) -> ToolResult:
    if not (directory / "manifest.json").exists():
        return ToolResult(
            status=ResultStatus.MISSING_SOURCE,
            warnings=["No registered topology source."],
            elapsed_ms=0,
        )
    if at is not None and at.utcoffset() is None:
        raise ValueError("topology lookup requires explicit time offset")
    manifest = json.loads((directory / "manifest.json").read_text())
    if file_sha(directory / "edges.json") != manifest["edges_sha256"]:
        raise ValueError("topology archive hash mismatch")
    device = canonical(device_id, manifest["device_aliases"])
    if interface_type is not None:
        interface_type = manifest["interface_aliases"].get(interface_type, interface_type)
    edges, excluded = [], []
    for raw in json.loads((directory / "edges.json").read_text()):
        edge = TopologyEdge.model_validate_json(json.dumps(raw))
        if device not in (edge.device_id, edge.peer_device_id):
            continue
        if interface_type is not None and edge.interface_type != interface_type:
            continue
        reason = (
            "before_valid_from"
            if at and edge.valid_from and at < edge.valid_from
            else "at_or_after_valid_to"
            if at and edge.valid_to and at >= edge.valid_to
            else None
        )
        if reason:
            excluded.append({**raw, "exclusion_reason": reason})
            continue
        edges.append(raw)
    warnings = []
    if manifest["synthetic"]:
        warnings.append("Synthetic fixture only; does not describe the real network.")
    if excluded:
        warnings.append(
            "Some matching edges were excluded by the requested observation time;"
            " excluded_edges are not active adjacency."
        )
    if at is None:
        warnings.append(
            "No observation time selected; applicability at incident time is unverified."
        )
    if any(
        any(
            flag in edge["quality_flags"]
            for flag in ("validity_unknown", "valid_from_unknown", "valid_to_unknown")
        )
        for edge in edges
    ):
        warnings.append(
            "Topology validity incomplete/unknown; returned edges are candidate checks only."
        )
    if manifest["quarantined_rows"]:
        warnings.append("Some imported rows are quarantined; topology may be incomplete.")
    truncated = len(edges) > 1000
    excluded_truncated = len(excluded) > 1000
    if excluded_truncated:
        warnings.append("Excluded edge details capped at 1000; excluded_total is complete.")
    return ToolResult(
        status=ResultStatus.MISSING_SOURCE
        if not edges
        else ResultStatus.PARTIAL
        if warnings or truncated
        else ResultStatus.COMPLETED,
        data={
            "edges": edges[:1000],
            "excluded_edges": excluded[:1000],
            "excluded_total": len(excluded),
            "excluded_truncated": excluded_truncated,
            "total_matches": len(edges),
            "truncated": truncated,
            "source_sha256": manifest["source_sha256"],
            "generation": manifest["generation"],
        },
        evidence_ids=[edge["edge_id"] for edge in edges[:1000]],
        warnings=warnings,
        elapsed_ms=0,
    )
